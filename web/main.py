#!/usr/bin/env python3
"""FastAPI backend for 네이버 부동산 매물 수집기 web app."""

import asyncio
import io
import os
import re
import sys
import tempfile
import threading
import uuid as _uuid
from pathlib import Path

import openpyxl
from fastapi import FastAPI, File, Query, UploadFile, WebSocket, WebSocketDisconnect
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import FileResponse, JSONResponse, StreamingResponse
from fastapi.staticfiles import StaticFiles
from pydantic import BaseModel

ROOT = Path(__file__).parent.parent
sys.path.insert(0, str(ROOT))

from naver_land_scraper import (
    append_row,
    collect_articles_by_url_list,
    is_duplicate,
    load_or_create_workbook,
    parse_url,
    search_region_articles,
)
from . import relay

app = FastAPI(title="네이버 부동산 매물 수집기")

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_methods=["*"],
    allow_headers=["*"],
)

STATIC_DIR = Path(__file__).parent / "static"
app.mount("/static", StaticFiles(directory=str(STATIC_DIR)), name="static")

# 세션별 임시 Excel 파일 저장 디렉터리
SESSIONS_DIR = Path(tempfile.gettempdir()) / "naver_land_sessions"
SESSIONS_DIR.mkdir(exist_ok=True)

# session_id → xlsx Path
_sessions: dict[str, Path] = {}

# job_id → asyncio.Queue
_jobs: dict[str, asyncio.Queue] = {}


def _session_path(session_id: str) -> Path:
    return SESSIONS_DIR / f"{session_id}.xlsx"


def _get_session_file(session_id: str) -> Path:
    """세션 파일 경로 반환. 없으면 빈 워크북 생성."""
    path = _session_path(session_id)
    if session_id not in _sessions or not path.exists():
        wb, _ = load_or_create_workbook(path)
        wb.save(path)
        _sessions[session_id] = path
    return path


def _count_rows(path: Path) -> int:
    try:
        wb = openpyxl.load_workbook(path, read_only=True)
        ws = wb.active
        count = max(0, ws.max_row - 1)  # 헤더 제외
        wb.close()
        return count
    except Exception:
        return 0


# ── Session endpoints ─────────────────────────────────────────────────────────

@app.get("/")
async def root():
    return FileResponse(str(STATIC_DIR / "index.html"))


@app.post("/api/session/new")
async def new_session():
    """새 빈 세션 생성 — 페이지 로드 시 호출."""
    try:
        sid = str(_uuid.uuid4())
        _get_session_file(sid)
        return {"session_id": sid}
    except Exception as e:
        return JSONResponse({"error": str(e)}, status_code=500)


@app.post("/api/excel/upload")
async def upload_excel(
    session_id: str = Query(...),
    file: UploadFile = File(...),
):
    """기존 Excel 파일 업로드 → 해당 세션의 파일로 사용."""
    if not (file.filename or "").endswith(".xlsx"):
        return JSONResponse({"error": ".xlsx 파일만 업로드 가능합니다"}, status_code=400)
    try:
        content = await file.read()
        wb = openpyxl.load_workbook(io.BytesIO(content))
        wb.close()
    except Exception as e:
        return JSONResponse({"error": f"유효하지 않은 Excel 파일입니다: {e}"}, status_code=400)

    try:
        path = _session_path(session_id)
        path.write_bytes(content)
        _sessions[session_id] = path
        rows = _count_rows(path)
        return {"filename": file.filename, "rows": rows}
    except Exception as e:
        return JSONResponse({"error": f"파일 저장 실패: {e}"}, status_code=500)


# ── WebSocket ─────────────────────────────────────────────────────────────────

@app.websocket("/ws/{job_id}")
async def ws_logs(websocket: WebSocket, job_id: str):
    """수집 진행 로그 스트림.

    아이폰 쪽 네트워크가 잠깐 끊겼다가 같은 job_id로 재연결하는 경우를 위해,
    작업이 실제로 끝나기(done) 전에는 큐를 지우지 않는다 — 재연결하면
    이어서 로그를 받을 수 있다.
    """
    await websocket.accept()
    q = _jobs.get(job_id)
    if q is None:
        await websocket.close(code=4004)
        return
    finished = False
    try:
        while True:
            try:
                msg = await asyncio.wait_for(q.get(), timeout=30)
            except asyncio.TimeoutError:
                try:
                    await websocket.send_json({"type": "ping"})
                except Exception:
                    break
                continue
            try:
                await websocket.send_json(msg)
            except Exception:
                await q.put(msg)  # 전송 실패 — 재연결 시 다시 받을 수 있도록 되돌려 놓음
                break
            if msg.get("type") == "done":
                finished = True
                break
    except WebSocketDisconnect:
        pass
    finally:
        if finished:
            _jobs.pop(job_id, None)


@app.websocket("/ws/relay/{session_id}")
async def ws_relay(websocket: WebSocket, session_id: str):
    """아이폰 앱이 붙어서, 이 서버 대신 네이버로 나가는 트래픽을 중계해준다.

    네이버가 클라우드 서버 IP를 차단하기 때문에, 실제 네이버 접속은 항상
    이 웹소켓에 연결된 아이폰을 거쳐서 나간다 (web/relay.py 참고).
    """
    await relay.phone_relay_endpoint(websocket, session_id)


# ── Request models ────────────────────────────────────────────────────────────

class RegionRequest(BaseModel):
    session_id: str
    regions: list[str]
    max_count: int = 500
    filters: dict = {}


class UrlRequest(BaseModel):
    session_id: str
    urls: list[str]
    filters: dict = {}


# ── Filter logic ──────────────────────────────────────────────────────────────

def _parse_price(s) -> float:
    """'13억5,000' / '13억 5,000' / '80,000' → 만원 단위 float. 파싱 불가 시 0."""
    s = str(s or "").replace(",", "").replace(" ", "")
    m = re.match(r"(\d+(?:\.\d+)?)억(\d*)", s)
    if m:
        return float(m.group(1)) * 10000 + (int(m.group(2)) if m.group(2) else 0)
    try:
        return float(s)
    except ValueError:
        return 0.0


def _parse_area(s) -> float | None:
    """'24.3평' → 24.3. 파싱 불가 시 None."""
    s = str(s or "").replace("평", "").replace(",", "").strip()
    try:
        return float(s)
    except ValueError:
        return None


def _floor_type(floor_s: str) -> str | None:
    """'5/15' → 저층/중층/고층. 이미 텍스트로 되어 있으면 그대로 반환."""
    m = re.match(r"(\d+)/(\d+)", floor_s or "")
    if m:
        cur, total = int(m.group(1)), int(m.group(2))
        if total <= 0:
            return None
        ratio = cur / total
        if ratio <= 0.33:
            return "저층"
        if ratio <= 0.66:
            return "중층"
        return "고층"
    for label in ("저층", "중층", "고층"):
        if label in (floor_s or ""):
            return label
    return None


def _passes_filter(fields: dict, filters: dict) -> bool:
    if not filters:
        return True

    price_min = filters.get("price_min")
    price_max = filters.get("price_max")
    area_min  = filters.get("area_min")
    area_max  = filters.get("area_max")
    directions = filters.get("directions", [])
    floors    = filters.get("floors", [])
    household_min = filters.get("household_min")

    if price_min or price_max:
        price = _parse_price(fields.get("price_main"))
        if price > 0:
            if price_min and price < float(price_min): return False
            if price_max and price > float(price_max): return False

    if area_min or area_max:
        area = _parse_area(fields.get("area_exclusive"))
        if area is not None:
            if area_min and area < float(area_min): return False
            if area_max and area > float(area_max): return False

    if directions:
        d = fields.get("direction", "")
        if d and d not in directions:
            return False

    if floors:
        ftype = _floor_type(fields.get("floor", ""))
        if ftype and ftype not in floors:
            return False

    if household_min:
        # household_by_type 형식: "300세대/25세대" 또는 "300세대"
        household_str = str(fields.get("household_by_type", ""))
        total_str = household_str.split("세대")[0].replace(",", "")
        try:
            if int(total_str) < int(household_min):
                return False
        except (ValueError, TypeError):
            pass

    return True


def _make_log_fn(q: asyncio.Queue, loop: asyncio.AbstractEventLoop):
    def log(msg: str, tag: str = "dim"):
        loop.call_soon_threadsafe(q.put_nowait, {"type": "log", "msg": msg, "tag": tag})
    return log


# ── Scraping endpoints ────────────────────────────────────────────────────────

@app.post("/api/scrape/region")
async def scrape_region(req: RegionRequest):
    if not relay.is_phone_connected(req.session_id):
        return JSONResponse(
            {"error": "아이폰 앱이 서버에 연결되어 있지 않습니다. 앱을 켜둔 상태로 다시 시도해 주세요."},
            status_code=409,
        )

    excel_path = _get_session_file(req.session_id)
    socks_port = await relay.get_or_start_socks_server(req.session_id)
    proxy = {"server": f"socks5://127.0.0.1:{socks_port}"}

    job_id = str(_uuid.uuid4())[:8]
    q: asyncio.Queue = asyncio.Queue()
    _jobs[job_id] = q
    loop = asyncio.get_event_loop()

    def run():
        total_ok = total_skip = total_filtered = 0
        for i, region in enumerate(req.regions, 1):
            log = _make_log_fn(q, loop)
            log(f"\n── [{i}/{len(req.regions)}] {region} ──", "accent")
            try:
                articles = search_region_articles(region, log=log, max_count=req.max_count, proxy=proxy)
            except RuntimeError as e:
                log(f"  ✗ 실패: {e}", "error"); continue
            except Exception as e:
                log(f"  ✗ 오류: {e}", "error"); continue

            try:
                wb, ws = load_or_create_workbook(excel_path)
            except Exception as e:
                log(f"  ✗ Excel 오류: {e}", "error"); continue

            ok = skip = filtered = 0
            for fields in articles:
                article_no = str(fields.get("article_no", ""))
                if not article_no: continue
                if is_duplicate(ws, article_no, fields.get("complex_name", ""), fields.get("floor", "")):
                    skip += 1; continue
                if not _passes_filter(fields, req.filters):
                    filtered += 1; continue
                append_row(ws, fields, ws.max_row + 1)
                ok += 1

            try:
                wb.save(excel_path)
            except Exception as e:
                log(f"  ✗ 저장 오류: {e}", "error"); continue

            log(f"  저장 {ok}건 / 중복 {skip}건 / 필터 {filtered}건", "info")
            total_ok += ok; total_skip += skip; total_filtered += filtered

        loop.call_soon_threadsafe(q.put_nowait, {
            "type": "done", "ok": total_ok, "skip": total_skip, "filtered": total_filtered,
            "rows": _count_rows(excel_path),
        })

    threading.Thread(target=run, daemon=True).start()
    return {"job_id": job_id}


@app.post("/api/scrape/urls")
async def scrape_urls(req: UrlRequest):
    if not relay.is_phone_connected(req.session_id):
        return JSONResponse(
            {"error": "아이폰 앱이 서버에 연결되어 있지 않습니다. 앱을 켜둔 상태로 다시 시도해 주세요."},
            status_code=409,
        )

    excel_path = _get_session_file(req.session_id)
    socks_port = await relay.get_or_start_socks_server(req.session_id)
    proxy = {"server": f"socks5://127.0.0.1:{socks_port}"}

    job_id = str(_uuid.uuid4())[:8]
    q: asyncio.Queue = asyncio.Queue()
    _jobs[job_id] = q
    loop = asyncio.get_event_loop()

    def run():
        log = _make_log_fn(q, loop)
        log(f"\n── URL 수집 ({len(req.urls)}개) ──", "accent")

        try:
            wb, ws = load_or_create_workbook(excel_path)
        except Exception as e:
            log(f"✗ Excel 오류: {e}", "error")
            loop.call_soon_threadsafe(q.put_nowait, {"type": "done", "ok": 0, "skip": 0, "filtered": 0, "rows": 0})
            return

        urls_to_fetch, skip = [], 0
        for url in req.urls:
            url = url.strip()
            if not url: continue
            try:
                article_no, _ = parse_url(url)
                if is_duplicate(ws, article_no):
                    log(f"  → 이미 등록된 매물 — 건너뜀 ({article_no})", "dim"); skip += 1
                else:
                    urls_to_fetch.append(url)
            except ValueError as e:
                log(f"  ✗ URL 오류: {e}", "error")

        if not urls_to_fetch:
            loop.call_soon_threadsafe(q.put_nowait, {"type": "done", "ok": 0, "skip": skip, "filtered": 0, "rows": _count_rows(excel_path)})
            return

        try:
            articles = collect_articles_by_url_list(urls_to_fetch, log=log, proxy=proxy)
        except Exception as e:
            log(f"✗ 수집 오류: {e}", "error")
            loop.call_soon_threadsafe(q.put_nowait, {"type": "done", "ok": 0, "skip": skip, "filtered": 0, "rows": _count_rows(excel_path)})
            return

        ok = filtered = 0
        for fields in articles:
            article_no = str(fields.get("article_no", ""))
            if is_duplicate(ws, article_no, fields.get("complex_name", ""), fields.get("floor", "")):
                skip += 1; continue
            if not _passes_filter(fields, req.filters):
                filtered += 1; continue
            append_row(ws, fields, ws.max_row + 1)
            ok += 1

        try:
            wb.save(excel_path)
        except Exception as e:
            log(f"✗ 저장 오류: {e}", "error")

        log(f"  저장 {ok}건 / 중복 {skip}건 / 필터 {filtered}건", "info")
        loop.call_soon_threadsafe(q.put_nowait, {
            "type": "done", "ok": ok, "skip": skip, "filtered": filtered,
            "rows": _count_rows(excel_path),
        })

    threading.Thread(target=run, daemon=True).start()
    return {"job_id": job_id}


# ── File endpoints ────────────────────────────────────────────────────────────

@app.get("/api/excel/download")
async def download_excel(session_id: str):
    path = _get_session_file(session_id)
    if not path.exists():
        return JSONResponse({"error": "파일이 없습니다"}, status_code=404)
    return FileResponse(
        str(path),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        filename="네이버_부동산_매물.xlsx",
    )


@app.post("/api/session/reset")
async def reset_session(session_id: str):
    """세션 파일을 빈 워크북으로 초기화."""
    path = _session_path(session_id)
    try:
        path.unlink(missing_ok=True)
    except Exception:
        pass
    _sessions.pop(session_id, None)
    _get_session_file(session_id)  # 빈 파일 재생성
    return {"ok": True}


@app.get("/api/excel/info")
async def excel_info(session_id: str):
    path = _get_session_file(session_id)
    return {"rows": _count_rows(path)}


@app.get("/api/excel/preview")
async def preview_excel(session_id: str, limit: int = 500):
    path = _get_session_file(session_id)
    try:
        wb = openpyxl.load_workbook(path, read_only=True)
        ws = wb.active
        headers = [str(cell.value or "") for cell in next(ws.iter_rows(min_row=1, max_row=1))]
        rows = []
        for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
            vals = [str(cell.value) if cell.value is not None else "" for cell in row]
            if any(vals):
                rows.append(vals)
            if len(rows) >= limit:
                break
        wb.close()
        return {"headers": headers, "rows": rows, "total": len(rows)}
    except StopIteration:
        return {"headers": [], "rows": [], "total": 0}
    except Exception as e:
        return JSONResponse({"error": str(e)}, status_code=500)


@app.post("/api/pdf/export")
async def export_pdf(session_id: str):
    path = _get_session_file(session_id)
    if _count_rows(path) == 0:
        return JSONResponse({"error": "저장된 매물이 없습니다"}, status_code=404)

    pdf_path = path.with_suffix(".pdf")
    try:
        import xlwings as xw
        xl_app = xw.App(visible=False)
        try:
            wb = xl_app.books.open(str(path))
            for sheet in wb.sheets:
                ps = sheet.api.page_setup
                ps.orientation = 2
                ps.fit_to_pages_wide = 1
                ps.fit_to_pages_tall = False
                pts = xl_app.api.inches_to_points
                ps.left_margin  = pts(0.25)
                ps.right_margin = pts(0.25)
                ps.top_margin   = pts(0.5)
                ps.bottom_margin = pts(0.5)
            wb.to_pdf(str(pdf_path))
            wb.close()
        finally:
            xl_app.quit()
        return FileResponse(str(pdf_path), media_type="application/pdf", filename="네이버_부동산_매물.pdf")
    except ImportError:
        return JSONResponse({"error": "xlwings가 설치되지 않았습니다"}, status_code=500)
    except Exception as e:
        return JSONResponse({"error": str(e)}, status_code=500)


if __name__ == "__main__":
    import uvicorn
    uvicorn.run("main:app", host="0.0.0.0", port=8000, reload=True)
